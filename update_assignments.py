from pathlib import Path
import shutil
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn

SRC = Path(r"D:\Duan_48\Bao_cao_hoan_chinh_KT1_De_tai_48.docx")
OUT = Path(r"D:\Duan_48\Bao_cao_hoan_chinh_KT1_De_tai_48_phan_cong.docx")

def set_cell(cell, text, size=11, bold=False):
    cell.text = text
    for p in cell.paragraphs:
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        for run in p.runs:
            run.font.name = "Times New Roman"
            run._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
            run._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
            run.font.size = Pt(size)
            run.bold = bold

def build():
    shutil.copy2(SRC, OUT)
    doc = Document(OUT)
    detailed = [
        ("1", "Tiêu chí 1: Phân tích bối cảnh, mục tiêu, phạm vi MVP và vấn đề quản lý", "Phạm Đức Việt"),
        ("2", "Tiêu chí 4: Actor, Use Case và quan hệ include/extend", "Phạm Đức Việt"),
        ("3", "Tiêu chí 6 & 10: Kiến trúc hệ thống, kế hoạch triển khai và tổng hợp báo cáo", "Phạm Đức Việt"),
        ("4", "Tiêu chí 2 & 3: Yêu cầu chức năng, phi chức năng và tiêu chí nghiệm thu", "Nguyễn Văn Hiếu"),
        ("5", "Tiêu chí 5: ERD, CSDL 9 bảng, PK/FK, ràng buộc và chỉ mục", "Nguyễn Văn Hiếu"),
        ("6", "Kiểm thử dữ liệu, kiểm soát rủi ro và rà soát công thức hóa đơn/công nợ", "Nguyễn Văn Hiếu"),
        ("7", "Luồng nghiệp vụ, dữ liệu đầu vào/đầu ra từ nhập chỉ số đến thanh toán", "Hoàng Văn Thoại"),
        ("8", "Tiêu chí 7 & 8: Vị trí AI, system/user prompt, JSON và luồng gọi AI", "Hoàng Văn Thoại"),
        ("9", "Tiêu chí 9: Minh chứng AI, test dữ liệu âm/mâu thuẫn và hậu kiểm phản hồi", "Hoàng Văn Thoại"),
    ]
    task_table = doc.tables[0]
    for row, values in zip(task_table.rows[1:], detailed):
        for cell, value in zip(row.cells, values):
            set_cell(cell, value, size=10.5)

    member_table = doc.tables[1]
    members = [
        ("1", "Phạm Đức Việt", "Trưởng nhóm; phân tích nghiệp vụ, actor/use case, kiến trúc, kế hoạch và tổng hợp báo cáo", ""),
        ("2", "Nguyễn Văn Hiếu", "Yêu cầu chức năng/phi chức năng, CSDL 9 bảng, kiểm thử và kiểm soát rủi ro", ""),
        ("3", "Hoàng Văn Thoại", "Luồng nghiệp vụ, dữ liệu đầu vào/đầu ra, AI prompt/JSON và minh chứng AI", ""),
    ]
    for row, values in zip(member_table.rows[1:], members):
        for cell, value in zip(row.cells, values):
            set_cell(cell, value, size=10.5)
    doc.save(OUT)
    print(OUT)

if __name__ == "__main__":
    build()
