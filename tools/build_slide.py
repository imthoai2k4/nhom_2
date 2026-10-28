"""Tạo slide báo cáo cuối kỳ UtilityAI (12 slide, 16:9)."""
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "UtilityAI_Slide.pptx"

W, H = 13.333, 7.5
BG = RGBColor(10, 15, 24)
PANEL = RGBColor(18, 26, 40)
TEXT = RGBColor(232, 238, 247)
MUTED = RGBColor(159, 176, 199)
GREEN = RGBColor(52, 211, 153)
YELLOW = RGBColor(251, 191, 36)
BLUE = RGBColor(96, 165, 250)
RED = RGBColor(248, 113, 113)
LINE = RGBColor(45, 58, 78)

FONT = "Segoe UI"


def rgb(hex_str):
    return RGBColor.from_string(hex_str)


def set_bg(slide, color=BG):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def textbox(slide, text, x, y, w, h, size=20, color=TEXT, bold=False,
            align=PP_ALIGN.LEFT, font=FONT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def bullets(slide, items, x, y, w, h, size=18, color=TEXT, gap=8):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        p.line_spacing = 1.12
        run = p.add_run()
        run.text = "•  " + item
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.color.rgb = color
    return box


def card(slide, value, label, x, y, w, h, accent=GREEN):
    shape = slide.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = PANEL
    shape.line.color.rgb = LINE
    shape.line.width = Pt(0.75)
    shape.shadow.inherit = False

    textbox(slide, value, x + 0.25, y + 0.18, w - 0.5, 0.6, 26, accent, True)
    textbox(slide, label, x + 0.25, y + 0.78, w - 0.5, h - 0.9, 12, MUTED)


def header(slide, num, kicker, title):
    set_bg(slide)
    # accent bar + slide number
    bar = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(W), Inches(0.06))
    bar.fill.solid(); bar.fill.fore_color.rgb = GREEN
    bar.line.fill.background()

    textbox(slide, kicker.upper(), 0.7, 0.35, 8.5, 0.3, 11, GREEN, True)
    textbox(slide, title, 0.7, 0.7, 11.8, 0.8, 30, TEXT, True)
    textbox(slide, f"{num:02d}", 12.2, 0.35, 0.7, 0.4, 12, MUTED, True, PP_ALIGN.RIGHT)
    textbox(slide, "UtilityAI · Đề tài 48", 0.7, 7.15, 6, 0.25, 9, MUTED)


def image_fit(slide, path, x, y, max_w, max_h):
    path = Path(path)
    if not path.exists():
        return None
    with Image.open(path) as im:
        iw, ih = im.size
    scale = min(max_w / iw, max_h / ih)
    w, h = iw * scale, ih * scale
    pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
    pic.line.color.rgb = LINE
    pic.line.width = Pt(0.75)
    return pic


def add_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


prs = Presentation()
prs.slide_width = Inches(W)
prs.slide_height = Inches(H)
blank = prs.slide_layouts[6]


# ── 1. Title ────────────────────────────────────────────────────────────────
s = prs.slides.add_slide(blank)
set_bg(s)
# decorative shapes
for cx, cy, d, color in [(10.6, 1.0, 2.2, PANEL), (11.8, 4.6, 1.4, PANEL), (9.6, 5.4, 0.8, PANEL)]:
    c = s.shapes.add_shape(9, Inches(cx), Inches(cy), Inches(d), Inches(d))
    c.fill.solid(); c.fill.fore_color.rgb = color
    c.line.color.rgb = LINE

textbox(s, "UTILITYAI", 0.9, 1.5, 6, 0.4, 14, GREEN, True)
textbox(s, "Quản lý Hóa đơn\nĐiện nước Hộ gia đình", 0.9, 1.95, 8.6, 1.9, 44, TEXT, True)
textbox(s, "Tích hợp AI phân tích tiêu thụ · Qwen3.8-27B", 0.9, 3.95, 8.5, 0.5, 20, MUTED)
textbox(s, "Nhóm 2  ·  Đề tài 48  ·  2026", 0.9, 6.2, 7, 0.4, 14, MUTED)
for i, label in enumerate(["FastAPI", "SQLite", "OpenAI", "94 tests"]):
    card(s, label, ["Backend", "Dữ liệu", "API chuẩn", "Kiểm thử"][i],
         0.9 + i * 2.25, 5.05, 2.0, 0.95, GREEN if i in (0, 2) else YELLOW)
add_notes(s, "Giới thiệu đề tài, mục tiêu demo 5-7 phút: quản lý điện nước + AI phân tích.")


# ── 2. Bài toán & mục tiêu ──────────────────────────────────────────────────
s = prs.slides.add_slide(blank)
header(s, 2, "Bối cảnh", "Bài toán & mục tiêu")
bullets(s, [
    "Hộ gia đình / khu trọ ghi chép chỉ số rời rạc, dễ sai và khó phát hiện bất thường.",
    "Cần quản lý hộ, đồng hồ, chỉ số theo kỳ, đơn giá, hóa đơn, thanh toán và công nợ.",
    "Muốn AI đọc dữ liệu thật để nhận xét, cảnh báo và gợi ý tiết kiệm — không bịa số.",
], 0.8, 1.9, 7.4, 2.2, 18)
card(s, "2", "vai trò: Admin / User", 0.8, 4.5, 3.6, 1.15, GREEN)
card(s, "8", "nhóm chức năng quản lý", 4.7, 4.5, 3.5, 1.15, YELLOW)
card(s, "3", "nhóm chức năng AI", 8.6, 4.5, 3.5, 1.15, BLUE)
image_fit(s, ROOT / "docs" / "ui-preview" / "admin-dashboard.png", 8.5, 1.65, 4.2, 2.55)
add_notes(s, "Nêu bài toán, phạm vi MVP và 2 vai trò. Ảnh dashboard minh họa giao diện.")


# ── 3. Kiến trúc ────────────────────────────────────────────────────────────
s = prs.slides.add_slide(blank)
header(s, 3, "Thiết kế", "Kiến trúc hệ thống")

def box(label, sub, x, y, w, h, accent=GREEN):
    sh = s.shapes.add_shape(5, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = PANEL
    sh.line.color.rgb = accent; sh.line.width = Pt(1)
    textbox(s, label, x + 0.12, y + 0.12, w - 0.24, 0.35, 14, TEXT, True)
    textbox(s, sub, x + 0.12, y + 0.52, w - 0.24, h - 0.6, 10, MUTED)

box("SPA", "HTML · CSS · Vanilla JS", 0.7, 1.9, 2.6, 1.0, BLUE)
box("FastAPI", "REST · Auth · RBAC", 4.0, 1.9, 2.6, 1.0, GREEN)
box("Services", "billing · alert · AI", 7.3, 1.9, 2.6, 1.0, YELLOW)
box("SQLite", "WAL · backup", 10.6, 1.9, 2.0, 1.0, BLUE)
box("AI Client", "OpenAI-compatible", 4.0, 4.2, 2.6, 1.0, GREEN)
box("Qwen3.8-27B", "Intern-AI · /chat/completions", 7.3, 4.2, 2.6, 1.0, YELLOW)
box("Prompt v3", "system + user + guard", 10.6, 4.2, 2.0, 1.0, YELLOW)

for x1, x2, y1, y2 in [(3.3, 4.0, 2.4, 2.4), (6.6, 7.3, 2.4, 2.4), (9.9, 10.6, 2.4, 2.4),
                       (5.3, 5.3, 2.9, 4.2), (5.3, 7.3, 4.7, 4.7)]:
    ln = s.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    ln.line.color.rgb = MUTED; ln.line.width = Pt(1.2)

textbox(s, "Dữ liệu đi từ CSDL → service → AI; kết quả quay lại UI kèm metadata nguồn/kỳ/prompt.",
        0.8, 6.1, 11.6, 0.4, 13, MUTED)
add_notes(s, "Giải thích luồng: SPA gọi API, service đọc SQLite, AI client gọi Qwen, prompt tách file.")


# ── 4. Chức năng quản lý ────────────────────────────────────────────────────
s = prs.slides.add_slide(blank)
header(s, 4, "Nghiệp vụ", "Quản lý hộ, chỉ số, hóa đơn, thanh toán")
bullets(s, [
    "Đăng nhập + phân quyền Admin/User theo hộ.",
    "CRUD hộ/phòng, đồng hồ, đơn giá bậc thang EVN + nước.",
    "Nhập chỉ số kỳ, kiểm tra trùng kỳ và chỉ số âm/giảm.",
    "Tính hóa đơn có VAT, thanh toán một phần/toàn bộ, công nợ.",
    "Cảnh báo tiêu thụ tăng > 50% so với TB 3 kỳ.",
], 0.8, 1.9, 6.4, 3.2, 17)
card(s, "4", "hộ demo", 0.8, 5.3, 3.0, 1.0, GREEN)
card(s, "7", "bậc giá điện", 4.0, 5.3, 3.0, 1.0, YELLOW)
image_fit(s, ROOT / "docs" / "ui-preview" / "user-readings.png", 7.5, 1.7, 5.2, 3.65)
add_notes(s, "Demo nhập chỉ số, tạo hóa đơn, ghi thanh toán, xem cảnh báo.")


# ── 5. AI pipeline ─────────────────────────────────────────────────────────
s = prs.slides.add_slide(blank)
header(s, 5, "AI", "Pipeline phân tích tiêu thụ")
bullets(s, [
    "Chọn N kỳ gần nhất (1–12) + hóa đơn + cảnh báo chưa xử lý.",
    "Prompt v3 tách system/user, nêu ngưỡng 50%, buộc 5 mục output.",
    "Gọi Qwen3.8-27B qua API chuẩn OpenAI; timeout 30s, max 4096 token.",
    "Kiểm tra đủ 5 mục + kết thúc câu; retry/fallback nếu thiếu.",
    "Trả metadata: source, model, kỳ, prompt version, thời điểm, cache.",
], 0.8, 1.9, 7.0, 3.3, 17)
card(s, "5", "mục output", 0.8, 5.35, 2.4, 1.0, GREEN)
card(s, "4096", "max tokens", 3.5, 5.35, 2.4, 1.0, YELLOW)
card(s, "30s", "timeout", 6.2, 5.35, 2.2, 1.0, BLUE)
image_fit(s, ROOT / "docs" / "ui-preview" / "admin-meters.png", 8.4, 1.75, 4.3, 3.0)
add_notes(s, "Nhấn mạnh prompt tách khỏi code, kiểm tra output và fallback nội bộ.")


# ── 6. Kết quả AI live ─────────────────────────────────────────────────────
s = prs.slides.add_slide(blank)
header(s, 6, "Kết quả", "AI trả lời đủ 5 mục, bám dữ liệu")
quote = ("“Tháng 7/2026, Phòng 101 dùng điện 78,00 kWh; tăng 3,00 kWh so với tháng 6 và tháng 5; "
         "nguyên nhân cụ thể: chưa có dữ liệu.”")
box = s.shapes.add_shape(5, Inches(0.8), Inches(1.9), Inches(11.7), Inches(1.5))
box.fill.solid(); box.fill.fore_color.rgb = PANEL; box.line.color.rgb = GREEN
textbox(s, quote, 1.1, 2.08, 11.1, 1.1, 19, TEXT)
for i, (v, l, c) in enumerate([
    ("1.035", "ký tự", GREEN), ("5/5", "mục", GREEN),
    ("12,1s", "thời gian", YELLOW), ("2.003", "token", BLUE),
]):
    card(s, v, l, 0.8 + i * 3.0, 3.9, 2.6, 1.0, c)
textbox(s, "Fallback nội bộ khi thiếu key / 429 / timeout / output không đạt — UI đổi badge nguồn.",
        0.8, 5.4, 11.6, 0.5, 15, MUTED)
add_notes(s, "Đây là kết quả live với Qwen3.8-27B; nhấn mạnh không bịa nguyên nhân.")


# ── 7. Chat AI ─────────────────────────────────────────────────────────────
s = prs.slides.add_slide(blank)
header(s, 7, "Chat", "Chat AI: tự nhiên nhưng có kiểm soát")
bullets(s, [
    "Câu hỏi tự do → gọi Qwen; tra cứu có cấu trúc → rule-based.",
    "Context chỉ chứa dữ liệu hộ được phép xem; admin xem hộ được hỏi.",
    "System prompt chặn đổi vai trò / lộ prompt / bịa số.",
    "Không có key hoặc lỗi → fallback hướng dẫn, không crash.",
], 0.8, 1.9, 6.0, 2.8, 17)
image_fit(s, ROOT / "docs" / "ui-preview" / "chat-window.png", 7.0, 1.55, 5.7, 4.9)
add_notes(s, "Demo: hỏi 'Theo bạn tháng 7 nhà tôi dùng điện thế nào?', sau đó thử prompt injection.")


# ── 8. Kiểm thử ────────────────────────────────────────────────────────────
s = prs.slides.add_slide(blank)
header(s, 8, "Chất lượng", "Kiểm thử: 94 passed")
rows = [
    ("AI + prompt + fallback", "32"),
    ("Chat AI", "7"),
    ("CRUD / công thức", "26"),
    ("API AI + bảo mật/CORS", "13"),
    ("Billing / Alert", "11"),
    ("Security / hash / rate limit", "5"),
]
y = 1.85
for name, count in rows:
    box = s.shapes.add_shape(1, Inches(0.8), Inches(y), Inches(7.6), Inches(0.55))
    box.fill.solid(); box.fill.fore_color.rgb = PANEL; box.line.color.rgb = LINE
    textbox(s, name, 1.0, y + 0.12, 5.8, 0.3, 13, TEXT)
    textbox(s, count, 7.4, y + 0.1, 0.8, 0.3, 14, GREEN, True, PP_ALIGN.RIGHT)
    y += 0.65
card(s, "94", "test cases", 9.2, 1.85, 3.2, 1.2, GREEN)
card(s, "0", "failed", 9.2, 3.25, 3.2, 1.2, BLUE)
textbox(s, "Bao gồm đúng/sai/biên: months 0/13, 429, timeout, output cắt, DESC, RBAC, cache, hash.",
        9.2, 4.8, 3.5, 1.2, 12, MUTED)
add_notes(s, "Nêu 94 test, các regression quan trọng: truncation, DESC, chat AI, RBAC, cache.")


# ── 9. Bảo mật & vận hành ──────────────────────────────────────────────────
s = prs.slides.add_slide(blank)
header(s, 9, "Vận hành", "Bảo mật & triển khai")
bullets(s, [
    "Mật khẩu PBKDF2-SHA256, tự nâng cấp plaintext cũ khi đăng nhập.",
    "CORS allowlist, rate limit AI/chat, security headers, HSTS khi HTTPS.",
    "SQLite WAL, backup WAL-safe, seed demo tách production.",
    "Nginx + systemd/NSSM, healthcheck /api/health, log AI latency/token.",
], 0.8, 1.9, 7.2, 3.0, 17)
card(s, "PBKDF2", "hash mật khẩu", 0.8, 4.9, 3.0, 1.0, GREEN)
card(s, "CORS", "allowlist", 4.0, 4.9, 3.0, 1.0, YELLOW)
card(s, "WAL", "SQLite", 7.2, 4.9, 3.0, 1.0, BLUE)
textbox(s, "deploy/nginx.conf · deploy/utilityai.service · tools/backup_db.py",
        0.8, 6.35, 11.5, 0.35, 12, MUTED)
add_notes(s, "Nêu các biện pháp bảo mật cơ bản và checklist deploy.")


# ── 10. Hạn chế ────────────────────────────────────────────────────────────
s = prs.slides.add_slide(blank)
header(s, 10, "Trung thực", "Hạn chế & hướng phát triển")
bullets(s, [
    "Auth demo dùng header username/password → production nên JWT/session + HTTPS.",
    "SQLite phù hợp demo nhỏ; nhiều người dùng nên PostgreSQL.",
    "AI có thể thiếu dữ liệu hoặc dừng giữa mục → đã có retry/fallback và cảnh báo UI.",
    "Chưa có mobile app, chưa tích hợp cổng thanh toán trực tuyến.",
], 0.8, 1.9, 7.3, 3.3, 17)
card(s, "→", "JWT/session", 0.8, 5.25, 3.2, 1.1, GREEN)
card(s, "→", "PostgreSQL", 4.2, 5.25, 3.2, 1.1, YELLOW)
card(s, "→", "thanh toán online", 7.6, 5.25, 3.6, 1.1, BLUE)
add_notes(s, "Nêu hạn chế thẳng thắn; nhấn mạnh fallback giúp demo ổn định.")


# ── 11. Demo flow ──────────────────────────────────────────────────────────
s = prs.slides.add_slide(blank)
header(s, 11, "Demo", "Kịch bản 5–7 phút")
steps = [
    ("1", "Admin đăng nhập → Dashboard"),
    ("2", "Nhập chỉ số → tạo hóa đơn → ghi thanh toán"),
    ("3", "Mở cảnh báo > 50%"),
    ("4", "Phân tích AI 6 tháng → xem 5 mục, badge, cache"),
    ("5", "Chat AI: hỏi tự do + thử prompt injection"),
    ("6", "Đăng nhập user → chỉ thấy dữ liệu hộ mình"),
]
y = 1.85
for num, step in steps:
    circ = s.shapes.add_shape(9, Inches(0.9), Inches(y), Inches(0.5), Inches(0.5))
    circ.fill.solid(); circ.fill.fore_color.rgb = GREEN
    circ.line.fill.background()
    textbox(s, num, 0.9, y + 0.1, 0.5, 0.3, 12, BG, True, PP_ALIGN.CENTER)
    textbox(s, step, 1.65, y + 0.08, 7.5, 0.4, 15, TEXT)
    y += 0.72
card(s, "429?", "dùng cache/fallback", 9.6, 2.3, 3.0, 1.0, YELLOW)
card(s, "mất mạng?", "phân tích nội bộ", 9.6, 3.6, 3.0, 1.0, BLUE)
add_notes(s, "Tập dượt 2 lần; phương án dự phòng 429/mất mạng là cache + fallback + ảnh chụp.")


# ── 12. Kết thúc ───────────────────────────────────────────────────────────
s = prs.slides.add_slide(blank)
set_bg(s)
textbox(s, "CẢM ƠN", 0.9, 1.5, 4, 0.4, 14, GREEN, True)
textbox(s, "Demo & Q&A", 0.9, 1.95, 8, 1.0, 44, TEXT, True)
textbox(s, "UtilityAI · Quản lý hóa đơn điện nước tích hợp AI", 0.9, 3.2, 9, 0.5, 18, MUTED)
textbox(s, "github.com/imthoai2k4/nhom_2", 0.9, 4.2, 8, 0.4, 14, BLUE)
textbox(s, "94 test · Qwen3.8-27B · FastAPI + SQLite", 0.9, 6.2, 8, 0.4, 13, MUTED)
add_notes(s, "Kết thúc, mời Q&A; mở repo/Swagger nếu cần.")

prs.save(OUT)
print(f"Saved {OUT} with {len(prs.slides)} slides")
